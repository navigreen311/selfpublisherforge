"""Report generation engine.

Generates PDF and XLSX reports for revenue summaries, book performance,
marketing ROI, and portfolio overviews.
"""

from __future__ import annotations

import io
import logging
import os
import uuid
from datetime import UTC, datetime
from typing import Any

from sqlalchemy import and_, func, select
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.ext.asyncio import AsyncSession

logger = logging.getLogger(__name__)

from app.modules.analytics.models import Report, RoyaltyRecord
from app.modules.analytics.schemas import OutputFormat, ReportStatus, ReportType


class ReportGenerationError(Exception):
    """Raised when report generation fails.

    Wraps underlying errors (database, I/O, library) that occur during
    PDF or XLSX report creation so callers receive a single, predictable
    exception type with the original cause available via ``__cause__``.
    """

    def __init__(self, message: str = "Report generation failed") -> None:
        self.message = message
        super().__init__(self.message)


async def generate_report(
    db: AsyncSession,
    report: Report,
) -> Report:
    """Generate a report based on its type and parameters.

    Updates the Report record with file path, size, and status.
    """
    try:
        report.status = ReportStatus.PROCESSING.value
        await db.flush()

        report_data = await _gather_report_data(
            db,
            org_id=report.org_id,
            report_type=ReportType(report.report_type),
            parameters=report.parameters or {},
        )

        output_format = OutputFormat(report.output_format)
        if output_format == OutputFormat.XLSX:
            file_bytes = _build_xlsx(report.title, report_data)
            extension = "xlsx"
        else:
            file_bytes = _build_pdf(report.title, report_data)
            extension = "pdf"

        # Store file to local disk (in production, upload to S3)
        reports_dir = os.path.join(os.getcwd(), "generated_reports")
        os.makedirs(reports_dir, exist_ok=True)
        file_name = f"{report.id}.{extension}"
        file_path = os.path.join(reports_dir, file_name)

        with open(file_path, "wb") as f:
            f.write(file_bytes)

        report.file_path = file_path
        report.file_size = len(file_bytes)
        report.status = ReportStatus.COMPLETED.value
        report.generated_at = datetime.now(UTC)
        await db.flush()

        return report

    except (ValueError, KeyError, OSError, SQLAlchemyError) as exc:
        logger.error("Report generation failed for report %s: %s", report.id, exc, exc_info=True)
        report.status = ReportStatus.FAILED.value
        report.error_message = str(exc)
        await db.flush()
        raise ReportGenerationError(str(exc)) from exc
    except Exception as exc:
        logger.exception("Unexpected error generating report %s", report.id, exc_info=True)
        report.status = ReportStatus.FAILED.value
        report.error_message = str(exc)
        await db.flush()
        raise ReportGenerationError(str(exc)) from exc


async def _gather_report_data(
    db: AsyncSession,
    org_id: uuid.UUID,
    report_type: ReportType,
    parameters: dict[str, Any],
) -> dict[str, Any]:
    """Gather data needed for a specific report type."""
    if report_type == ReportType.REVENUE_SUMMARY:
        return await _gather_revenue_summary(db, org_id, parameters)
    if report_type == ReportType.BOOK_PERFORMANCE:
        return await _gather_book_performance(db, org_id, parameters)
    if report_type == ReportType.PORTFOLIO_OVERVIEW:
        return await _gather_portfolio_overview(db, org_id, parameters)
    if report_type == ReportType.MARKETING_ROI:
        return await _gather_marketing_roi(db, org_id, parameters)
    return await _gather_revenue_summary(db, org_id, parameters)


async def _gather_revenue_summary(
    db: AsyncSession,
    org_id: uuid.UUID,
    parameters: dict[str, Any],
) -> dict[str, Any]:
    """Gather data for a revenue summary report."""
    # Total revenue
    totals_query = (
        select(
            func.coalesce(func.sum(RoyaltyRecord.net_revenue), 0).label("total_revenue"),
            func.coalesce(func.sum(RoyaltyRecord.gross_revenue), 0).label("gross_revenue"),
            func.coalesce(func.sum(RoyaltyRecord.net_units), 0).label("total_units"),
            func.count(RoyaltyRecord.id).label("record_count"),
        )
        .where(
            and_(
                RoyaltyRecord.org_id == org_id,
                RoyaltyRecord.deleted_at.is_(None),
            )
        )
    )
    result = await db.execute(totals_query)
    totals = result.one()

    # Revenue by platform
    platform_query = (
        select(
            RoyaltyRecord.platform,
            func.coalesce(func.sum(RoyaltyRecord.net_revenue), 0).label("revenue"),
            func.coalesce(func.sum(RoyaltyRecord.net_units), 0).label("units"),
        )
        .where(
            and_(
                RoyaltyRecord.org_id == org_id,
                RoyaltyRecord.deleted_at.is_(None),
            )
        )
        .group_by(RoyaltyRecord.platform)
    )
    platform_result = await db.execute(platform_query)
    platforms = [
        {"platform": r.platform, "revenue": str(r.revenue), "units": r.units}
        for r in platform_result.all()
    ]

    return {
        "report_type": "Revenue Summary",
        "total_revenue": str(totals.total_revenue),
        "gross_revenue": str(totals.gross_revenue),
        "total_units": totals.total_units,
        "record_count": totals.record_count,
        "by_platform": platforms,
    }


async def _gather_book_performance(
    db: AsyncSession,
    org_id: uuid.UUID,
    parameters: dict[str, Any],
) -> dict[str, Any]:
    """Gather data for a book performance report."""
    books_query = (
        select(
            RoyaltyRecord.title,
            func.coalesce(func.sum(RoyaltyRecord.net_revenue), 0).label("revenue"),
            func.coalesce(func.sum(RoyaltyRecord.net_units), 0).label("units"),
            func.count(RoyaltyRecord.id).label("records"),
        )
        .where(
            and_(
                RoyaltyRecord.org_id == org_id,
                RoyaltyRecord.deleted_at.is_(None),
            )
        )
        .group_by(RoyaltyRecord.title)
        .order_by(func.sum(RoyaltyRecord.net_revenue).desc())
        .limit(50)
    )
    result = await db.execute(books_query)
    books = [
        {"title": r.title, "revenue": str(r.revenue), "units": r.units, "records": r.records}
        for r in result.all()
    ]

    return {
        "report_type": "Book Performance",
        "books": books,
        "total_books": len(books),
    }


async def _gather_portfolio_overview(
    db: AsyncSession,
    org_id: uuid.UUID,
    parameters: dict[str, Any],
) -> dict[str, Any]:
    """Gather data for a portfolio overview report."""
    revenue_data = await _gather_revenue_summary(db, org_id, parameters)
    book_data = await _gather_book_performance(db, org_id, parameters)
    return {
        "report_type": "Portfolio Overview",
        **revenue_data,
        **book_data,
    }


async def _gather_marketing_roi(
    db: AsyncSession,
    org_id: uuid.UUID,
    parameters: dict[str, Any],
) -> dict[str, Any]:
    """Gather data for a marketing ROI report.

    Combines revenue data with advertising spend from the campaigns module
    to compute return on investment.
    """
    from app.modules.advertising.models import Campaign, CampaignPerformance

    revenue_data = await _gather_revenue_summary(db, org_id, parameters)

    # Compute total advertising spend
    spend_stmt = (
        select(func.coalesce(func.sum(CampaignPerformance.spend), 0.0))
        .join(Campaign, CampaignPerformance.campaign_id == Campaign.id)
        .where(
            Campaign.org_id == org_id,
            Campaign.deleted_at.is_(None),
            CampaignPerformance.deleted_at.is_(None),
        )
    )
    spend_result = await db.execute(spend_stmt)
    total_spend = float(spend_result.scalar() or 0.0)
    total_revenue = float(revenue_data.get("total_revenue", 0))
    roi = ((total_revenue - total_spend) / total_spend * 100) if total_spend > 0 else 0.0

    return {
        "report_type": "Marketing ROI",
        "total_ad_spend": round(total_spend, 2),
        "total_revenue": round(total_revenue, 2),
        "net_return": round(total_revenue - total_spend, 2),
        "roi_percent": round(roi, 1),
        **revenue_data,
    }


def _build_pdf(title: str, data: dict[str, Any]) -> bytes:
    """Build a PDF report using ReportLab.

    Generates a formatted PDF with title page header, summary metrics,
    and data tables for each report section.  Returns raw PDF bytes.
    """
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import inch
        from reportlab.platypus import (
            Paragraph,
            SimpleDocTemplate,
            Spacer,
            Table,
            TableStyle,
        )
    except ImportError as exc:
        raise ReportGenerationError(
            "The 'reportlab' package is required for PDF generation. "
            "Install it with:  pip install reportlab"
        ) from exc

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=letter,
        leftMargin=0.75 * inch,
        rightMargin=0.75 * inch,
        topMargin=0.75 * inch,
        bottomMargin=0.75 * inch,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ReportTitle",
        parent=styles["Title"],
        fontSize=22,
        spaceAfter=6,
        textColor=colors.HexColor("#1a1a2e"),
    )
    subtitle_style = ParagraphStyle(
        "ReportSubtitle",
        parent=styles["Normal"],
        fontSize=10,
        textColor=colors.grey,
        spaceAfter=20,
    )
    section_style = ParagraphStyle(
        "SectionHeading",
        parent=styles["Heading2"],
        fontSize=14,
        spaceBefore=18,
        spaceAfter=8,
        textColor=colors.HexColor("#16213e"),
    )
    elements: list[Any] = []

    # --- Header ---
    generated_at = datetime.now(UTC).strftime("%Y-%m-%d %H:%M UTC")
    report_type = data.get("report_type", "Report")

    elements.append(Paragraph(title, title_style))
    elements.append(
        Paragraph(f"{report_type}  |  Generated: {generated_at}", subtitle_style)
    )

    # --- Horizontal rule via thin table ---
    hr = Table([[""]],  colWidths=[7 * inch], rowHeights=[1])
    hr.setStyle(TableStyle([
        ("LINEABOVE", (0, 0), (-1, 0), 1, colors.HexColor("#0f3460")),
    ]))
    elements.append(hr)
    elements.append(Spacer(1, 12))

    # --- Summary Metrics ---
    summary_rows: list[list[str]] = []
    if "total_revenue" in data:
        summary_rows.append(["Total Revenue", f"${data['total_revenue']}"])
    if "gross_revenue" in data:
        summary_rows.append(["Gross Revenue", f"${data['gross_revenue']}"])
    if "total_units" in data:
        summary_rows.append(["Total Units Sold", str(data["total_units"])])
    if "record_count" in data:
        summary_rows.append(["Royalty Records", str(data["record_count"])])
    if "total_ad_spend" in data:
        summary_rows.append(["Total Ad Spend", f"${data['total_ad_spend']}"])
    if "net_return" in data:
        summary_rows.append(["Net Return", f"${data['net_return']}"])
    if "roi_percent" in data:
        summary_rows.append(["ROI", f"{data['roi_percent']}%"])
    if "total_books" in data:
        summary_rows.append(["Books Tracked", str(data["total_books"])])

    if summary_rows:
        elements.append(Paragraph("Summary", section_style))
        t = Table(summary_rows, colWidths=[3 * inch, 4 * inch])
        t.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 10),
            ("TEXTCOLOR", (0, 0), (0, -1), colors.HexColor("#16213e")),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("LINEBELOW", (0, 0), (-1, -2), 0.5, colors.lightgrey),
            ("LINEBELOW", (0, -1), (-1, -1), 1, colors.HexColor("#0f3460")),
        ]))
        elements.append(t)
        elements.append(Spacer(1, 12))

    # --- Platform Breakdown Table ---
    if data.get("by_platform"):
        elements.append(Paragraph("Revenue by Platform", section_style))
        header = ["Platform", "Revenue", "Units"]
        rows = [header] + [
            [p["platform"], f"${p['revenue']}", str(p["units"])]
            for p in data["by_platform"]
        ]
        t = Table(rows, colWidths=[2.5 * inch, 2.25 * inch, 2.25 * inch])
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f3460")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 10),
            ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f0f0f5")]),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.lightgrey),
        ]))
        elements.append(t)
        elements.append(Spacer(1, 12))

    # --- Book Performance Table ---
    if data.get("books"):
        elements.append(Paragraph("Book Performance", section_style))
        header = ["Title", "Revenue", "Units", "Records"]
        rows = [header] + [
            [
                b["title"][:50] + ("..." if len(b["title"]) > 50 else ""),
                f"${b['revenue']}",
                str(b["units"]),
                str(b.get("records", "")),
            ]
            for b in data["books"][:25]
        ]
        t = Table(rows, colWidths=[3 * inch, 1.5 * inch, 1.25 * inch, 1.25 * inch])
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f3460")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f0f0f5")]),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.lightgrey),
        ]))
        elements.append(t)

    # --- Footer note ---
    elements.append(Spacer(1, 24))
    footer_style = ParagraphStyle(
        "Footer",
        parent=styles["Normal"],
        fontSize=8,
        textColor=colors.grey,
    )
    elements.append(
        Paragraph(
            f"Self Publisher Forge  &bull;  Report generated {generated_at}",
            footer_style,
        )
    )

    doc.build(elements)
    return buf.getvalue()


def _build_xlsx(title: str, data: dict[str, Any]) -> bytes:
    """Build an XLSX workbook using openpyxl.

    Creates a workbook with a Summary sheet and additional data sheets for
    platform breakdown and book performance when the data is available.
    Returns raw XLSX bytes.
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side, numbers
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise ReportGenerationError(
            "The 'openpyxl' package is required for XLSX generation. "
            "Install it with:  pip install openpyxl"
        ) from exc

    wb = Workbook()

    # ---- Shared style helpers ----
    header_font = Font(name="Calibri", bold=True, size=11, color="FFFFFF")
    header_fill = PatternFill(start_color="0F3460", end_color="0F3460", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC"),
    )
    title_font = Font(name="Calibri", bold=True, size=16, color="1A1A2E")
    subtitle_font = Font(name="Calibri", size=10, color="888888")
    label_font = Font(name="Calibri", bold=True, size=11)
    currency_format = '#,##0.00'
    integer_format = '#,##0'

    def _style_header_row(ws: Any, row: int, max_col: int) -> None:
        """Apply header styling to a row."""
        for col_idx in range(1, max_col + 1):
            cell = ws.cell(row=row, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border

    def _auto_width(ws: Any, min_width: int = 12, max_width: int = 40) -> None:
        """Set column widths based on content length."""
        for col_cells in ws.columns:
            length = min_width
            col_letter = get_column_letter(col_cells[0].column)
            for cell in col_cells:
                if cell.value is not None:
                    length = max(length, min(len(str(cell.value)) + 2, max_width))
            ws.column_dimensions[col_letter].width = length

    generated_at = datetime.now(UTC).strftime("%Y-%m-%d %H:%M UTC")
    report_type = data.get("report_type", "Report")

    # ================================================================
    # SHEET 1 -- Summary
    # ================================================================
    ws_summary = wb.active
    ws_summary.title = "Summary"
    ws_summary.sheet_properties.tabColor = "0F3460"

    # Title block
    ws_summary.merge_cells("A1:D1")
    title_cell = ws_summary["A1"]
    title_cell.value = title
    title_cell.font = title_font

    ws_summary.merge_cells("A2:D2")
    sub_cell = ws_summary["A2"]
    sub_cell.value = f"{report_type}  |  Generated: {generated_at}"
    sub_cell.font = subtitle_font

    row = 4  # start summary metrics at row 4

    metric_pairs: list[tuple[str, str]] = []
    if "total_revenue" in data:
        metric_pairs.append(("Total Revenue", data["total_revenue"]))
    if "gross_revenue" in data:
        metric_pairs.append(("Gross Revenue", data["gross_revenue"]))
    if "total_units" in data:
        metric_pairs.append(("Total Units Sold", str(data["total_units"])))
    if "record_count" in data:
        metric_pairs.append(("Royalty Records", str(data["record_count"])))
    if "total_ad_spend" in data:
        metric_pairs.append(("Total Ad Spend", str(data["total_ad_spend"])))
    if "net_return" in data:
        metric_pairs.append(("Net Return", str(data["net_return"])))
    if "roi_percent" in data:
        metric_pairs.append(("ROI (%)", str(data["roi_percent"])))
    if "total_books" in data:
        metric_pairs.append(("Books Tracked", str(data["total_books"])))

    for label, value in metric_pairs:
        label_cell = ws_summary.cell(row=row, column=1, value=label)
        label_cell.font = label_font

        val_cell = ws_summary.cell(row=row, column=2)
        try:
            numeric = float(value)
            val_cell.value = numeric
            if label.lower() in ("total units sold", "royalty records", "books tracked"):
                val_cell.number_format = integer_format
            elif "roi" in label.lower():
                val_cell.number_format = '0.0"%"'
            else:
                val_cell.number_format = currency_format
        except (ValueError, TypeError):
            val_cell.value = value

        row += 1

    _auto_width(ws_summary)

    # ================================================================
    # SHEET 2 -- Platform Breakdown (if applicable)
    # ================================================================
    if data.get("by_platform"):
        ws_plat = wb.create_sheet("Platform Breakdown")
        ws_plat.sheet_properties.tabColor = "16213E"

        headers = ["Platform", "Revenue", "Units"]
        for col_idx, header in enumerate(headers, start=1):
            ws_plat.cell(row=1, column=col_idx, value=header)
        _style_header_row(ws_plat, 1, len(headers))

        for i, p in enumerate(data["by_platform"], start=2):
            ws_plat.cell(row=i, column=1, value=p["platform"]).border = thin_border

            rev_cell = ws_plat.cell(row=i, column=2)
            try:
                rev_cell.value = float(p["revenue"])
                rev_cell.number_format = currency_format
            except (ValueError, TypeError):
                rev_cell.value = p["revenue"]
            rev_cell.border = thin_border

            units_cell = ws_plat.cell(row=i, column=3)
            units_cell.value = p["units"]
            units_cell.number_format = integer_format
            units_cell.border = thin_border

        _auto_width(ws_plat)

    # ================================================================
    # SHEET 3 -- Book Performance (if applicable)
    # ================================================================
    if data.get("books"):
        ws_books = wb.create_sheet("Book Performance")
        ws_books.sheet_properties.tabColor = "E94560"

        headers = ["Title", "Revenue", "Units", "Records"]
        for col_idx, header in enumerate(headers, start=1):
            ws_books.cell(row=1, column=col_idx, value=header)
        _style_header_row(ws_books, 1, len(headers))

        for i, b in enumerate(data["books"], start=2):
            ws_books.cell(row=i, column=1, value=b["title"]).border = thin_border

            rev_cell = ws_books.cell(row=i, column=2)
            try:
                rev_cell.value = float(b["revenue"])
                rev_cell.number_format = currency_format
            except (ValueError, TypeError):
                rev_cell.value = b["revenue"]
            rev_cell.border = thin_border

            units_cell = ws_books.cell(row=i, column=3)
            units_cell.value = b["units"]
            units_cell.number_format = integer_format
            units_cell.border = thin_border

            rec_cell = ws_books.cell(row=i, column=4)
            rec_cell.value = b.get("records", "")
            rec_cell.border = thin_border

        # Title column wider
        ws_books.column_dimensions["A"].width = 45
        _auto_width(ws_books, min_width=14)
        ws_books.column_dimensions["A"].width = 45  # re-set after auto

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
